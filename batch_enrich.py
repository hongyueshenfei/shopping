#!/usr/bin/env python3
"""
Batch Enrich - 批量商品信息补全脚本

从 Excel 读取商品名称列表，调用 enrich_product() 逐条补全分类/品牌/品类/品种/百科信息，
将结果追加写入原文件的右侧新列（每次运行会覆盖之前的补全列）。

输出文件: <原文件名>_补全.xlsx（如 1年内销售库存大于10育种+上海_补全.xlsx）

使用方法:
    python3 batch_enrich.py

依赖:
    product_enricher.py（核心补全逻辑）
    商品分类目录.xlsx（分类参考目录）
    1年内销售库存大于10育种+上海.xlsx（源数据文件）
"""

import sys
import os
import json
import re
from pathlib import Path

# 将当前目录加入导入路径，以便引用 product_enricher 模块
sys.path.insert(0, str(Path(__file__).parent))

from product_enricher import (
    load_categories, enrich_product, get_api_key, get_base_url,
    build_category_tree
)

import anthropic
import urllib.request
import urllib.parse


# ===========================================================================
# 文件路径工具
# ===========================================================================

def get_excel_path():
    """
    查找源 Excel 文件。
    优先查找脚本同目录，其次查找用户桌面路径。
    返回 Path 对象，文件不存在返回 None。
    """
    candidates = [
        Path(__file__).parent / '1年内销售库存大于10育种+上海.xlsx',
        Path.home() / 'Desktop' / '沈飞mac' / 'skills' / 'shopping' / '1年内销售库存大于10育种+上海.xlsx',
    ]
    for p in candidates:
        if p.exists():
            return p
    return None


def get_output_path(input_path):
    """
    生成输出文件路径：原文件名 + "_补全" 后缀。
    例: xxx.xlsx → xxx_补全.xlsx
    """
    return input_path.parent / f"{input_path.stem}_补全{input_path.suffix}"


# ===========================================================================
# 主流程
# ===========================================================================

def main():
    # ---- 1. 定位源文件 ----
    excel_path = get_excel_path()
    if not excel_path:
        print("错误: 找不到 '1年内销售库存大于10育种+上海.xlsx'")
        sys.exit(1)

    import openpyxl
    print(f"读取文件: {excel_path}")
    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active

    total_rows = ws.max_row - 1  # 减去表头行
    print(f"共 {total_rows} 条记录待处理\n")

    # ---- 2. 加载分类目录 ----
    categories = None
    excel_dir = excel_path.parent
    for parent in [excel_dir, Path(__file__).parent]:
        xlsx = parent / '商品分类目录.xlsx'
        if xlsx.exists():
            print(f"加载分类目录: {xlsx}")
            categories = load_categories(str(xlsx))
            print(f"共 {len(categories)} 条分类记录\n")
            break

    if not categories:
        print("错误: 找不到商品分类目录.xlsx")
        sys.exit(1)

    # ---- 3. 定义输出列 ----
    # 原始列（保持不动）
    header_cols = [
        'goods_no', 'goods_name', 'department_name', 'department_code',
        'all_qty', 'all_amount', '结存数量'
    ]
    # 新增补全列
    enrich_cols = [
        '一级分类', '二级分类', '三级分类', '四级分类',
        '品牌', '品类', '品种', '百科信息'
    ]

    # 写入表头（覆盖已有表头，确保列名最新）
    for i, col_name in enumerate(header_cols + enrich_cols, 1):
        ws.cell(1, i).value = col_name

    # ---- 4. 逐行处理 ----
    success_count = 0
    error_count = 0

    for row_idx in range(2, ws.max_row + 1):
        goods_name = ws.cell(row_idx, 2).value  # 商品名称在第2列（B列）
        if not goods_name:
            continue

        # 进度条显示
        current = row_idx - 1
        pct = current / total_rows * 100
        bar_len = 30
        filled = int(bar_len * current / total_rows)
        bar = '█' * filled + '░' * (bar_len - filled)
        print(f"\r[{bar}] {pct:.1f}% ({current}/{total_rows}) - {goods_name[:20]}", end='', flush=True)

        # 调用核心补全函数
        result = enrich_product(str(goods_name), '', categories)

        # 写入结果到右侧新列
        if 'error' in result:
            error_count += 1
            # 失败时清空补全列
            for i, _ in enumerate(enrich_cols):
                ws.cell(row_idx, len(header_cols) + i + 1).value = ''
        else:
            success_count += 1
            # 成功时将8个字段分别写入独立列
            for i, col_name in enumerate(enrich_cols):
                val = result.get(col_name, '')
                ws.cell(row_idx, len(header_cols) + i + 1).value = val

        # 每50行保存一次（防止意外中断丢失数据）
        if (row_idx - 2) % 50 == 0:
            wb.save(get_output_path(excel_path))

    print(f"\n\n处理完成！成功: {success_count}, 失败: {error_count}")

    # ---- 5. 最终保存 ----
    out_path = get_output_path(excel_path)
    wb.save(out_path)
    print(f"结果已保存: {out_path}")


if __name__ == '__main__':
    main()